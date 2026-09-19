"""Generates the fictionalised inbound order-fulfilment documents for Corvina Quarter Stage 2:
  1. <PO>_Order_Schedule.xlsx - a call-up style order schedule (blanket PO, a sub-PO per delivery event, per-level delivery
     dates on two call-up schedules, grouped by delivery window), modelled on the layout of a real commercial call-up sheet.
     The programme (events, dates, windows) is the builder's programme in meridian-data-v4 files 21-22.
  2. <PO>_email.eml          - the customer's PO email that carries it.
All parties, addresses, numbers and contacts are fictional. Items are evolved from the tender schedule
(Corvina_Quarter_Stage2_Finishes_Schedule_RevC.xlsx); "ALT" marks an approved alternative, derived only from the
catalogue's Discontinued -> SupersededBy links (never from the Build Plan 2 answer key)."""
import csv, datetime as dt, os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.normpath(os.path.join(HERE, "..", "..", "meridian-data-v2"))
PO, BSO, PROJECT_ID = "BPO-0441", "BSO-2026-0141", "PID-CQ-S2"
OUT_XLSX = os.path.join(DATA, f"Corvina_Quarter_{PO}_Order_Schedule.xlsx")
OUT_EML = os.path.join(DATA, f"Corvina_Quarter_{PO}_email.eml")
TODAY = dt.date(2026, 9, 19)

products = {r["ModelCode"]: r for r in csv.DictReader(open(os.path.join(DATA, "04_products.csv"), encoding="utf-8-sig"))}
by_code = {r["ProductCode"]: r for r in products.values()}
families = {r["Name"]: r["Category"] for r in csv.DictReader(open(os.path.join(DATA, "02_product_families.csv"), encoding="utf-8-sig"))}

# ---- read tender schedule items ----
FAMILY_CODE = {"Wall Oven": "OVN", "Cooktop": "CT", "Dishwasher": "DW", "Rangehood": "RH", "Microwave": "MW", "Basin": "BSN",
               "Toilet Suite": "TS", "Basin Mixer": "MX", "Shower Set": "SS", "Kitchen Sink Mixer": "KM"}
ws = openpyxl.load_workbook(os.path.join(DATA, "Corvina_Quarter_Stage2_Finishes_Schedule_RevC.xlsx"), data_only=True).active
header_row = next(r for r in range(1, 20) if ws.cell(r, 1).value == "Item")
items, seq = [], {}
for r in range(header_row + 1, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, 14)]
    if not v[0] or not v[8]:
        continue
    model = str(v[5] or "").strip()
    prod = products.get(model)
    alt = bool(prod and prod["LifecycleStatus"] == "Discontinued" and prod["SupersededByCode"])
    if alt:
        prod = by_code[prod["SupersededByCode"]]
        model = prod["ModelCode"]
    family = prod["ProductFamily"] if prod else ""
    group = "Appliances" if (families.get(family) == "Appliances" or (not prod and v[1] in ("Kitchen", "Laundry", "Butlers Pantry"))) else "Tapware & Sanitaryware"
    code = FAMILY_CODE.get(family, "GEN")  # item ref: architect-style family code + sequence (DW-01, CT-03)
    seq[code] = seq.get(code, 0) + 1
    ref = f"{code}-{seq[code]:02d}"
    items.append(dict(item=f"{ref}{' ALT' if alt else ''}", tender_item=str(v[0]), model=model or "(unconfirmed)",
                      desc=str(v[3])[:60], tier=str(v[2]), qty=int(v[8]), group=group))

# ---- delivery programme: Corvina's call-up schedules (builder programme, v4), one sub-PO per delivery event ----
V4 = os.path.normpath(os.path.join(HERE, "..", "..", "meridian-data-v4"))
SCHEDULE = {"Appliances": "SCH-0141-01", "Tapware & Sanitaryware": "SCH-0141-02"}
windows = {w["Code"]: w["Name"] for w in csv.DictReader(open(os.path.join(V4, "22_delivery_windows.csv"), encoding="utf-8-sig"))}
events = [e for e in csv.DictReader(open(os.path.join(V4, "21_corvina_events.csv"), encoding="utf-8-sig")) if e["ScheduleRef"] in SCHEDULE.values()]
for n, e in enumerate(sorted(events, key=lambda e: (e["ScheduledOn"], e["EventCode"])), start=1):
    e["subpo"] = f"{PO}-{n:02d}"
    e["date"] = dt.date.fromisoformat(e["ScheduledOn"])
    e["window"] = windows[e["WindowCode"]]
by_group = {g: sorted([e for e in events if e["ScheduleRef"] == s], key=lambda e: int(e["Sequence"])) for g, s in SCHEDULE.items()}

def split(total, parts):
    base, extra = divmod(total, parts)
    return [base + (1 if i < extra else 0) for i in range(parts)]

for it in items:
    evs = by_group[it["group"]]
    levels = [e for e in evs if e["EventType"] == "Level rollout"]
    q = {e["EventCode"]: 0 for e in evs}
    remaining = it["qty"]
    proto = next((e for e in evs if e["EventType"] == "Prototype"), None)
    if proto and it["tier"] == "Standard" and remaining > len(levels):
        q[proto["EventCode"]] = 1; remaining -= 1
    if it["tier"] == "Penthouse":
        q[levels[-1]["EventCode"]] = remaining  # penthouse finishes land with the top level
    else:
        for e, n in zip(levels, split(remaining, len(levels))): q[e["EventCode"]] = n
    it["plan"] = q
    it["received"] = sum(n for c, n in q.items() if next(e for e in evs if e["EventCode"] == c)["date"] < TODAY)
    it["remaining"] = it["qty"] - it["received"]

# ---- workbook ----
wb = openpyxl.Workbook(); sh = wb.active; sh.title = "Order Schedule"
bold, thin = Font(bold=True), Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"Received": "C6EFCE", "Overdue": "FFC7CE", "Scheduled for delivery": "BDD7EE"}
for r, t in enumerate(["Meridian Commercial Supply Pty Ltd", "Unit 4, 88 Dohertys Road", "Truganina VIC 3029",
                       "T: 03 9000 0141", "ABN: 00 000 000 000 (fictional)"], start=1):
    sh.cell(r, 7, t).font = Font(bold=(r == 1), size=12 if r == 1 else 10)
sh.cell(7, 2, "APARTMENT FINISHES - ORDER SCHEDULE / CALL-UP").font = Font(bold=True, size=13)
info = [("Customer", "HALLORAN BRIGHT CONSTRUCTIONS PTY LTD (fictional)"), ("Blanket Sales Ref #", BSO), ("Customer PO No.", PO),
        ("Project ID/Name", f"{PROJECT_ID}  Corvina Quarter Stage 2")]
site = [("Delivery Address:", "12 Harbour Esplanade, Docklands VIC 3008"), ("Site Contact:", "Idris Fanshawe (Project Manager)"),
        ("Phone #:", "0400 000 141"), ("Email:", "idris.fanshawe@halloranbright.example")]
for i, (k, v) in enumerate(info):
    sh.cell(8 + i, 2, k).font = bold; sh.cell(8 + i, 3, v)
for i, (k, v) in enumerate(site):
    sh.cell(8 + i, 5, k).font = bold; sh.cell(8 + i, 6, v)
for i, (k, c) in enumerate(FILL.items()):
    sh.cell(7 + i, 16, "").fill = PatternFill("solid", fgColor=c); sh.cell(7 + i, 17, k)

FIRST = 9  # first delivery column
heads = ["Item Ref", "Tender Item #", "Model Code", "Item Description", "Order Qty", "Remaining Qty", f"Received on Site ({TODAY:%d/%m/%y})"]
row = 13
for grp in ("Appliances", "Tapware & Sanitaryware"):
    g_items = [i for i in items if i["group"] == grp]
    if not g_items:
        continue
    sh.cell(row, 2, f"{grp.upper()} - delivery schedule {SCHEDULE[grp]}").font = Font(bold=True, size=11); row += 1
    evs = by_group[grp]
    for c, d in enumerate(evs, start=FIRST):  # window / sub-PO / date / label header rows
        sh.cell(row, c, d["window"]).font = Font(italic=True, size=8)
        sh.cell(row + 1, c, d["subpo"]).font = Font(size=8)
        cell = sh.cell(row + 2, c, d["date"].strftime("%d/%m/%Y"))
        cell.fill = PatternFill("solid", fgColor=FILL["Received" if d["date"] < TODAY else "Scheduled for delivery"])
        sh.cell(row + 3, c, d["Label"]).font = bold
    sh.cell(row, FIRST - 1, "Delivery window").font = bold
    sh.cell(row + 1, FIRST - 1, "Sub-PO #").font = bold
    sh.cell(row + 2, FIRST - 1, "Delivery Date").font = bold
    for c, h in enumerate(heads, start=2):
        sh.cell(row + 3, c, h).font = bold
    row += 4
    for it in g_items:
        vals = [it["item"], it["tender_item"], it["model"], it["desc"], it["qty"], it["remaining"], it["received"]]
        for c, v in enumerate(vals, start=2):
            sh.cell(row, c, v).border = box
        for c, d in enumerate(evs, start=FIRST):
            n = it["plan"][d["EventCode"]]
            cell = sh.cell(row, c, n if n else None); cell.border = box
            cell.alignment = Alignment(horizontal="center")
        row += 1
    row += 2
sh.cell(row, 2, "Quantities are totals per delivery. Each delivery column is a separate sub-PO under the blanket PO above; "
                "items marked ALT are approved alternatives to the tendered model.").font = Font(italic=True, size=9)
sh.column_dimensions["E"].width = 44
for c in "BCDFGH": sh.column_dimensions[c].width = 14
wb.save(OUT_XLSX)

# ---- PO email ----
total = sum(i["qty"] for i in items)
eml = f"""From: Dana Whitlock <dana.whitlock@halloranbright.example>
To: Commercial Orders <orders@meridiancommercial.example>
Cc: Idris Fanshawe <idris.fanshawe@halloranbright.example>
Date: Fri, 18 Sep 2026 16:42:00 +1000
Subject: PO {PO} - Corvina Quarter Stage 2 - Appliances, tapware & sanitaryware (Blanket {BSO})
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"

Hi Meridian Commercial team,

Following award of tender TND-2026-0141, please find attached our blanket purchase order {PO}
for Corvina Quarter Stage 2 (project {PROJECT_ID}), referencing your blanket sales order {BSO}.

- {len(items)} line items, {total} units in total, across {len(events)} scheduled deliveries on two call-up
  schedules: appliances ({SCHEDULE["Appliances"]}) and plumbing and sanitaryware ({SCHEDULE["Tapware & Sanitaryware"]}).
- Each delivery column in the attached schedule is issued as its own sub-PO ({PO}-01 to {PO}-{len(events):02d}).
- The onsite prototype is due {min(e["date"] for e in events):%d/%m/%Y}; plumbing runs ahead of appliances, level by level
  through Level 24.
- Items marked ALT are the approved alternatives agreed at tender clarification.
- Deliveries to 12 Harbour Esplanade, Docklands. Site contact: Idris Fanshawe, 0400 000 141.

Please confirm stock and delivery dates against each sub-PO, and flag any line you cannot fill
from your Melbourne DC so we can agree store or interstate supply early.

Attachment: Corvina_Quarter_{PO}_Order_Schedule.xlsx

Kind regards,
Dana Whitlock
Contracts Manager, Halloran Bright Constructions (fictional)
"""
open(OUT_EML, "w").write(eml)
print(f"{len(items)} items ({sum(1 for i in items if 'ALT' in i['item'])} ALT), {len(events)} deliveries/sub-POs, {total} units")
print(OUT_XLSX); print(OUT_EML)
